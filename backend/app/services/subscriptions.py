"""Subscription cycles, institutional inventory, and premium-key redemption."""

from dataclasses import dataclass
from datetime import datetime, timezone
from calendar import monthrange
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.core.security import hash_password, verify_password
from app.models.ai import AiUserQuota, AiUserSubscription
from app.models.pre_registration import PreRegistration
from app.models.subscription import InstitutionAdminScope, InstitutionSubscriptionAllocation, PremiumSubscriptionKey
from app.models.user import MentorProfile, Role, StudentProfile, User
from app.services.admin_scope import can_manage_institution, institution_admin_abbrs, is_super_admin, user_institution_abbr
from app.services.spreadsheet_safety import safe_spreadsheet_text
from app.services.subscription_keys import generate_subscription_key, is_valid_subscription_key_format, normalize_subscription_key, subscription_key_fingerprint

PLAN_CREDITS = {"basic": 10, "pro": 50, "ultra": 100, "max": 200}
PREMIUM_PLANS = frozenset({"pro", "ultra", "max"})


class SubscriptionError(RuntimeError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class SubscriptionSnapshot:
    plan_code: str
    credit_limit: int
    credit_balance: int
    cycle_started_at: datetime
    cycle_ends_at: datetime

    @property
    def credits_used(self) -> int:
        return max(0, self.credit_limit - self.credit_balance)


@dataclass(frozen=True)
class InstitutionOption:
    institution_abbr: str
    institution_name_zh: str


@dataclass(frozen=True)
class IssuedSubscriptionKey:
    key: PremiumSubscriptionKey
    plaintext: str


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _as_utc(value: datetime) -> datetime:
    return value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)


def add_calendar_month(value: datetime) -> datetime:
    """Advance one calendar month, clamping 29--31 to the last valid day."""
    value = _as_utc(value)
    year = value.year + (value.month // 12)
    month = value.month % 12 + 1
    return value.replace(year=year, month=month, day=min(value.day, monthrange(year, month)[1]))


def _get_quota(db: Session, user_id, *, lock: bool = False) -> AiUserQuota | None:
    statement = select(AiUserQuota).where(AiUserQuota.user_id == user_id)
    if lock:
        statement = statement.with_for_update()
    return db.scalar(statement)


def _synchronize_quota(db: Session, subscription: AiUserSubscription, *, reset_balance: bool) -> AiUserQuota:
    quota = _get_quota(db, subscription.user_id, lock=True)
    if quota is None:
        quota = AiUserQuota(
            user_id=subscription.user_id,
            plan_code=subscription.plan_code,
            daily_limit=subscription.credit_limit,
            credit_balance=subscription.credit_limit,
        )
        db.add(quota)
        db.flush()
        return quota
    quota.plan_code = subscription.plan_code
    # Retained column name for a backward-compatible migration. It now represents
    # the current subscription-cycle cap, not a per-day allowance.
    quota.daily_limit = subscription.credit_limit
    if reset_balance:
        quota.credit_balance = subscription.credit_limit
    else:
        quota.credit_balance = min(quota.credit_balance, subscription.credit_limit)
    return quota


def ensure_user_subscription(db: Session, user: User, *, lock: bool = False, now: datetime | None = None) -> AiUserSubscription:
    statement = select(AiUserSubscription).where(AiUserSubscription.user_id == user.id)
    if lock:
        statement = statement.with_for_update()
    subscription = db.scalar(statement)
    if subscription is not None:
        return subscription
    start = _as_utc(user.created_at or now or utc_now())
    subscription = AiUserSubscription(
        user_id=user.id,
        plan_code="basic",
        cycle_started_at=start,
        cycle_ends_at=add_calendar_month(start),
        credit_limit=PLAN_CREDITS["basic"],
    )
    db.add(subscription)
    db.flush()
    _synchronize_quota(db, subscription, reset_balance=True)
    return subscription


def refresh_subscription(db: Session, user: User, *, lock: bool = False, now: datetime | None = None) -> AiUserSubscription:
    """Apply calendar-cycle resets and premium-to-basic fallback lazily and atomically."""
    current = _as_utc(now or utc_now())
    subscription = ensure_user_subscription(db, user, lock=lock, now=current)
    changed = False
    while current >= _as_utc(subscription.cycle_ends_at):
        next_start = _as_utc(subscription.cycle_ends_at)
        # A premium entitlement is only valid for the cycle activated by its key.
        # The first cycle after it expires is always the normal institution plan.
        if subscription.plan_code in PREMIUM_PLANS:
            subscription.plan_code = "basic"
        subscription.cycle_started_at = next_start
        subscription.cycle_ends_at = add_calendar_month(next_start)
        subscription.credit_limit = PLAN_CREDITS[subscription.plan_code]
        changed = True
    _synchronize_quota(db, subscription, reset_balance=changed)
    return subscription


def subscription_snapshot(db: Session, user: User, *, lock: bool = False) -> SubscriptionSnapshot:
    subscription = refresh_subscription(db, user, lock=lock)
    quota = _synchronize_quota(db, subscription, reset_balance=False)
    return SubscriptionSnapshot(
        plan_code=subscription.plan_code,
        credit_limit=subscription.credit_limit,
        credit_balance=quota.credit_balance,
        cycle_started_at=subscription.cycle_started_at,
        cycle_ends_at=subscription.cycle_ends_at,
    )


def active_premium_subscription_user_ids(
    db: Session,
    user_ids: list[object],
    *,
    now: datetime | None = None,
) -> set[object]:
    """Return users whose current premium cycle has not yet ended.

    Expired premium records are deliberately excluded. The normal activation
    path resets them to the basic plan before redemption.
    """
    if not user_ids:
        return set()
    current = _as_utc(now or utc_now())
    return set(
        db.scalars(
            select(AiUserSubscription.user_id).where(
                AiUserSubscription.user_id.in_(user_ids),
                AiUserSubscription.plan_code.in_(PREMIUM_PLANS),
                AiUserSubscription.cycle_ends_at > current,
            )
        ).all()
    )


def _allocation_field(plan_code: str) -> str:
    if plan_code not in PREMIUM_PLANS:
        raise SubscriptionError("invalid_subscription_plan")
    return f"{plan_code}_credits"


def _allocation(db: Session, tenant_id, institution_abbr: str, *, lock: bool = False) -> InstitutionSubscriptionAllocation | None:
    statement = select(InstitutionSubscriptionAllocation).where(
        InstitutionSubscriptionAllocation.tenant_id == tenant_id,
        InstitutionSubscriptionAllocation.institution_abbr == institution_abbr,
    )
    if lock:
        statement = statement.with_for_update()
    return db.scalar(statement)


def set_institution_allocation(
    db: Session,
    admin: User,
    *,
    institution_abbr: str,
    pro_credits: int,
    ultra_credits: int,
    max_credits: int,
) -> InstitutionSubscriptionAllocation:
    if not is_super_admin(admin):
        raise SubscriptionError("super_admin_required")
    institution = _institution_option(db, admin, institution_abbr)
    allocation = _allocation(db, admin.tenant_id, institution.institution_abbr, lock=True)
    if allocation is None:
        allocation = InstitutionSubscriptionAllocation(
            tenant_id=admin.tenant_id,
            institution_abbr=institution.institution_abbr,
            institution_name_zh=institution.institution_name_zh,
        )
        db.add(allocation)
    allocation.institution_name_zh = institution.institution_name_zh
    allocation.pro_credits = pro_credits
    allocation.ultra_credits = ultra_credits
    allocation.max_credits = max_credits
    db.commit()
    db.refresh(allocation)
    return allocation


def list_allocations(db: Session, admin: User) -> list[InstitutionSubscriptionAllocation]:
    statement = select(InstitutionSubscriptionAllocation).where(InstitutionSubscriptionAllocation.tenant_id == admin.tenant_id)
    if not is_super_admin(admin):
        scopes = institution_admin_abbrs(db, admin)
        statement = statement.where(InstitutionSubscriptionAllocation.institution_abbr.in_(scopes))
    return list(db.scalars(statement.order_by(InstitutionSubscriptionAllocation.institution_abbr)).all())


def list_institutions(db: Session, super_admin: User) -> list[InstitutionOption]:
    """Return trusted institutions already represented by imported or active accounts."""
    if not is_super_admin(super_admin):
        raise SubscriptionError("super_admin_required")
    institutions: dict[str, str] = {}
    sources = (
        select(PreRegistration.institution_abbr, PreRegistration.institution_name_zh).where(
            PreRegistration.tenant_id == super_admin.tenant_id
        ),
        select(StudentProfile.institution_abbr, StudentProfile.university)
        .join(User, User.id == StudentProfile.user_id)
        .where(User.tenant_id == super_admin.tenant_id),
        select(MentorProfile.institution_abbr, MentorProfile.university)
        .join(User, User.id == MentorProfile.user_id)
        .where(User.tenant_id == super_admin.tenant_id),
    )
    for statement in sources:
        for abbr, name in db.execute(statement).all():
            normalized_abbr = (abbr or "").strip().upper()
            normalized_name = (name or "").strip()
            if normalized_abbr and normalized_name:
                institutions.setdefault(normalized_abbr, normalized_name)
    return [InstitutionOption(institution_abbr=abbr, institution_name_zh=name) for abbr, name in sorted(institutions.items())]


def _institution_option(db: Session, super_admin: User, institution_abbr: str) -> InstitutionOption:
    normalized = institution_abbr.strip().upper()
    for option in list_institutions(db, super_admin):
        if option.institution_abbr == normalized:
            return option
    raise SubscriptionError("institution_not_found")


def list_institution_accounts(db: Session, admin: User, institution_abbr: str) -> list[User]:
    """Return only active verified student/mentor accounts in an authorized school."""
    normalized_abbr = institution_abbr.strip().upper()
    if not can_manage_institution(db, admin, normalized_abbr):
        raise SubscriptionError("institution_scope_forbidden")
    # Super administrators additionally require a trusted/imported institution;
    # scoped institution administrators may enumerate only their own scope.
    if is_super_admin(admin):
        normalized_abbr = _institution_option(db, admin, normalized_abbr).institution_abbr
    statement = (
        select(User)
        .options(selectinload(User.roles), selectinload(User.student_profile), selectinload(User.mentor_profile))
        .where(
            User.tenant_id == admin.tenant_id,
            User.is_active.is_(True),
            User.is_verified.is_(True),
            or_(
                User.student_profile.has(StudentProfile.institution_abbr == normalized_abbr),
                User.mentor_profile.has(MentorProfile.institution_abbr == normalized_abbr),
            ),
        )
        .order_by(User.full_name, User.username)
    )
    return list(db.scalars(statement).unique().all())


def assign_institution_admin(
    db: Session,
    super_admin: User,
    *,
    user_id,
    institution_abbr: str,
) -> InstitutionAdminScope:
    if not is_super_admin(super_admin):
        raise SubscriptionError("super_admin_required")
    institution = _institution_option(db, super_admin, institution_abbr)
    target = db.scalar(
        select(User).options(selectinload(User.roles), selectinload(User.student_profile), selectinload(User.mentor_profile)).where(
            User.tenant_id == super_admin.tenant_id,
            User.id == user_id,
        )
    )
    if (
        target is None
        or not target.is_active
        or not target.is_verified
        or user_institution_abbr(target) != institution.institution_abbr
    ):
        raise SubscriptionError("institution_admin_user_ineligible")
    role = db.scalar(
        select(Role).where(Role.tenant_id == super_admin.tenant_id, Role.code == "institution_admin")
    )
    if role is None:
        raise SubscriptionError("institution_admin_role_unavailable")
    if not any(item.code == role.code for item in target.roles):
        target.roles.append(role)
    scope = db.scalar(
        select(InstitutionAdminScope)
        .where(
            InstitutionAdminScope.tenant_id == super_admin.tenant_id,
            InstitutionAdminScope.user_id == target.id,
            InstitutionAdminScope.institution_abbr == institution.institution_abbr,
        )
        .with_for_update()
    )
    if scope is None:
        scope = InstitutionAdminScope(
            tenant_id=super_admin.tenant_id,
            user_id=target.id,
            institution_abbr=institution.institution_abbr,
            institution_name_zh=institution.institution_name_zh,
            assigned_by_user_id=super_admin.id,
        )
        db.add(scope)
    else:
        scope.institution_name_zh = institution.institution_name_zh
    db.commit()
    db.refresh(scope)
    return scope


def list_institution_admin_scopes(db: Session, admin: User) -> list[tuple[InstitutionAdminScope, User]]:
    if not is_super_admin(admin):
        raise SubscriptionError("super_admin_required")
    rows = db.execute(
        select(InstitutionAdminScope, User)
        .join(User, User.id == InstitutionAdminScope.user_id)
        .where(InstitutionAdminScope.tenant_id == admin.tenant_id)
        .order_by(InstitutionAdminScope.institution_abbr, User.username)
    ).all()
    return list(rows)


def list_my_institution_admin_scopes(db: Session, admin: User) -> list[InstitutionAdminScope]:
    return list(
        db.scalars(
            select(InstitutionAdminScope)
            .where(InstitutionAdminScope.tenant_id == admin.tenant_id, InstitutionAdminScope.user_id == admin.id)
            .order_by(InstitutionAdminScope.institution_abbr)
        ).all()
    )


def revoke_institution_admin(db: Session, super_admin: User, scope_id) -> None:
    if not is_super_admin(super_admin):
        raise SubscriptionError("super_admin_required")
    scope = db.scalar(
        select(InstitutionAdminScope)
        .where(InstitutionAdminScope.id == scope_id, InstitutionAdminScope.tenant_id == super_admin.tenant_id)
        .with_for_update()
    )
    if scope is None:
        raise SubscriptionError("institution_admin_scope_not_found")
    target = db.scalar(
        select(User).options(selectinload(User.roles)).where(User.id == scope.user_id, User.tenant_id == super_admin.tenant_id).with_for_update()
    )
    db.delete(scope)
    db.flush()
    has_remaining_scope = db.scalar(
        select(InstitutionAdminScope.id)
        .where(InstitutionAdminScope.tenant_id == super_admin.tenant_id, InstitutionAdminScope.user_id == scope.user_id)
        .limit(1)
    )
    if target is not None and has_remaining_scope is None:
        target.roles[:] = [role for role in target.roles if role.code != "institution_admin"]
    db.commit()


def issue_subscription_keys(
    db: Session,
    admin: User,
    *,
    institution_abbr: str,
    plan_code: str,
    quantity: int,
) -> list[IssuedSubscriptionKey]:
    if not can_manage_institution(db, admin, institution_abbr):
        raise SubscriptionError("institution_scope_forbidden")
    allocation = _allocation(db, admin.tenant_id, institution_abbr, lock=True)
    if allocation is None:
        raise SubscriptionError("institution_allocation_not_found")
    field = _allocation_field(plan_code)
    available = getattr(allocation, field)
    if quantity < 1 or quantity > available:
        raise SubscriptionError("subscription_inventory_exhausted")
    issued: list[IssuedSubscriptionKey] = []
    for _ in range(quantity):
        for _ in range(12):
            plaintext = generate_subscription_key(plan_code)
            fingerprint = subscription_key_fingerprint(plaintext)
            existing_key = db.scalar(select(PremiumSubscriptionKey.id).where(PremiumSubscriptionKey.key_fingerprint == fingerprint))
            existing_access_key = db.scalar(select(PreRegistration.id).where(PreRegistration.access_key_fingerprint == fingerprint))
            if existing_key is not None or existing_access_key is not None:
                continue
            key = PremiumSubscriptionKey(
                tenant_id=admin.tenant_id,
                institution_abbr=institution_abbr,
                institution_name_zh=allocation.institution_name_zh,
                plan_code=plan_code,
                key_hash=hash_password(plaintext),
                key_fingerprint=fingerprint,
                issued_by_user_id=admin.id,
            )
            try:
                # A savepoint keeps keys created earlier in the same batch intact
                # if a concurrent issuer wins the globally unique fingerprint race.
                with db.begin_nested():
                    db.add(key)
                    db.flush()
            except IntegrityError:
                continue
            issued.append(IssuedSubscriptionKey(key=key, plaintext=plaintext))
            break
        else:
            db.rollback()
            raise SubscriptionError("subscription_key_generation_failed")
    setattr(allocation, field, available - quantity)
    db.commit()
    for item in issued:
        db.refresh(item.key)
    return issued


def issue_subscription_key(db: Session, admin: User, *, institution_abbr: str, plan_code: str) -> tuple[PremiumSubscriptionKey, str]:
    """Compatibility helper for callers that issue exactly one key."""
    issued = issue_subscription_keys(db, admin, institution_abbr=institution_abbr, plan_code=plan_code, quantity=1)
    return issued[0].key, issued[0].plaintext


def _style_receipt_header(cells) -> None:
    for cell in cells:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build_subscription_key_receipt(issued: list[IssuedSubscriptionKey]) -> bytes:
    """Build the one-time Excel receipt; plaintext keys are never persisted."""
    if not issued:
        raise ValueError("At least one issued key is required")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订阅 Key 回执 Subscription Keys"
    sheet.append(["高级订阅 Key", "订阅档位", "院校中文全称", "院校简称", "签发时间"])
    _style_receipt_header(list(sheet[1]))
    for item in issued:
        sheet.append([
            safe_spreadsheet_text(item.plaintext),
            safe_spreadsheet_text(item.key.plan_code.upper()),
            safe_spreadsheet_text(item.key.institution_name_zh),
            safe_spreadsheet_text(item.key.institution_abbr),
            safe_spreadsheet_text(item.key.issued_at.isoformat() if item.key.issued_at else ""),
        ])
    for column, width in zip("ABCDE", (24, 16, 34, 18, 28), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    notice = workbook.create_sheet("安全说明 Security Notice")
    notice.append(["安全说明 / Security Notice"])
    _style_receipt_header([notice["A1"]])
    notice.append(["本回执包含仅此一次提供的高级订阅 Key 明文。请经受控渠道交付给所属院校的对应用户；系统不会保存 Key 明文。"])
    notice.append(["Key 激活后不可收回；未激活 Key 收回后永久失效。请妥善保管此 Excel 文件。"])
    notice.column_dimensions["A"].width = 120
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
    db.rollback()
    raise SubscriptionError("subscription_key_generation_failed")


def list_subscription_keys(db: Session, admin: User) -> list[PremiumSubscriptionKey]:
    statement = select(PremiumSubscriptionKey).where(PremiumSubscriptionKey.tenant_id == admin.tenant_id)
    if not is_super_admin(admin):
        scopes = institution_admin_abbrs(db, admin)
        statement = statement.where(PremiumSubscriptionKey.institution_abbr.in_(scopes))
    return list(db.scalars(statement.order_by(PremiumSubscriptionKey.issued_at.desc())).all())


def revoke_subscription_key(db: Session, admin: User, key_id) -> PremiumSubscriptionKey:
    key = db.scalar(
        select(PremiumSubscriptionKey).where(PremiumSubscriptionKey.id == key_id, PremiumSubscriptionKey.tenant_id == admin.tenant_id).with_for_update()
    )
    if key is None:
        raise SubscriptionError("subscription_key_not_found")
    if not can_manage_institution(db, admin, key.institution_abbr):
        raise SubscriptionError("institution_scope_forbidden")
    if key.status != "issued":
        raise SubscriptionError("subscription_key_not_reclaimable")
    allocation = _allocation(db, admin.tenant_id, key.institution_abbr, lock=True)
    if allocation is None:
        raise SubscriptionError("institution_allocation_not_found")
    setattr(allocation, _allocation_field(key.plan_code), getattr(allocation, _allocation_field(key.plan_code)) + 1)
    key.status = "revoked"
    key.revoked_at = utc_now()
    db.commit()
    db.refresh(key)
    return key


def _redeemable_subscription_key(db: Session, user: User, raw_key: str) -> PremiumSubscriptionKey:
    normalized = normalize_subscription_key(raw_key)
    if not is_valid_subscription_key_format(normalized):
        raise SubscriptionError("subscription_key_invalid")
    key = db.scalar(
        select(PremiumSubscriptionKey)
        .where(PremiumSubscriptionKey.key_fingerprint == subscription_key_fingerprint(normalized))
        .with_for_update()
    )
    # The response intentionally does not reveal whether a key ever existed.
    # Verify a dummy Argon2id hash for an unknown fingerprint too, preventing
    # a practical timing oracle that distinguishes an issued key from a typo.
    if not verify_password(normalized, key.key_hash if key is not None else None) or key is None or key.status != "issued":
        raise SubscriptionError("subscription_key_invalid")
    if key.tenant_id != user.tenant_id:
        raise SubscriptionError("subscription_key_invalid")
    db.refresh(user, attribute_names=["student_profile", "mentor_profile"])
    if user_institution_abbr(user) != key.institution_abbr:
        raise SubscriptionError("subscription_key_institution_mismatch")
    return key


def _activate_locked_subscription_key(
    db: Session,
    user: User,
    key: PremiumSubscriptionKey,
    *,
    started_at: datetime | None = None,
) -> SubscriptionSnapshot:
    """Consume an already locked, verified key without committing the session."""
    subscription = refresh_subscription(db, user, lock=True)
    if subscription.plan_code != "basic":
        raise SubscriptionError("active_premium_plan_locked")
    started_at = _as_utc(started_at or utc_now())
    subscription.plan_code = key.plan_code
    subscription.cycle_started_at = started_at
    subscription.cycle_ends_at = add_calendar_month(started_at)
    subscription.credit_limit = PLAN_CREDITS[key.plan_code]
    quota = _synchronize_quota(db, subscription, reset_balance=True)
    key.status = "activated"
    key.activated_by_user_id = user.id
    key.activated_at = started_at
    return SubscriptionSnapshot(
        plan_code=subscription.plan_code,
        credit_limit=subscription.credit_limit,
        credit_balance=quota.credit_balance,
        cycle_started_at=subscription.cycle_started_at,
        cycle_ends_at=subscription.cycle_ends_at,
    )


def activate_subscription_key(db: Session, user: User, *, raw_key: str) -> SubscriptionSnapshot:
    key = _redeemable_subscription_key(db, user, raw_key)
    snapshot = _activate_locked_subscription_key(db, user, key)
    db.commit()
    return snapshot
