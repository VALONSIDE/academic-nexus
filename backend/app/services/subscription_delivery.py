"""In-memory import, validation, activation, and offline PDF delivery of keys."""

from __future__ import annotations

import secrets
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.security import verify_password
from app.models.subscription import PremiumSubscriptionKey
from app.models.user import User
from app.services.admin_scope import can_manage_institution, user_institution_abbr
from app.services.subscription_delivery_pdf import render_subscription_delivery_pdf
from app.services.subscription_keys import is_valid_subscription_key_format, normalize_subscription_key, subscription_key_fingerprint
from app.services.subscriptions import (
    SubscriptionError,
    _activate_locked_subscription_key,
    active_premium_subscription_user_ids,
)


# System-issued receipts for up to 1,000 keys are far below this size.  Keeping
# the cap below Starlette's in-memory spool threshold avoids writing imported
# Key material to a temporary server file.
MAX_WORKBOOK_BYTES = 512 * 1024
MAX_WORKBOOK_UNCOMPRESSED_BYTES = 2 * 1024 * 1024
MAX_DELIVERY_ITEMS = 300


@dataclass(frozen=True)
class ImportedSubscriptionKey:
    row_number: int
    raw_key: str


@dataclass(frozen=True)
class DeliveryValidation:
    row_number: int
    raw_key: str
    plan_code: str | None
    institution_abbr: str | None
    institution_name_zh: str | None
    status: str


@dataclass(frozen=True)
class DeliveryItem:
    raw_key: str
    recipient_user_id: object


@dataclass(frozen=True)
class ResolvedDeliveryItem:
    raw_key: str
    key: PremiumSubscriptionKey
    recipient: User | None
    general_delivery: bool = False


def extract_subscription_keys_from_workbook(content: bytes) -> list[ImportedSubscriptionKey]:
    """Scan a single XLSX in memory.  Formulas are never evaluated or retained."""
    if not content or len(content) > MAX_WORKBOOK_BYTES:
        raise SubscriptionError("subscription_delivery_workbook_invalid")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 200 or sum(item.file_size for item in archive.infolist()) > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                raise SubscriptionError("subscription_delivery_workbook_invalid")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise SubscriptionError("subscription_delivery_workbook_invalid") from exc
    extracted: list[ImportedSubscriptionKey] = []
    seen: set[str] = set()
    try:
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                for value in row:
                    if not isinstance(value, str):
                        continue
                    normalized = normalize_subscription_key(value)
                    if normalized in seen or not is_valid_subscription_key_format(normalized):
                        continue
                    seen.add(normalized)
                    extracted.append(ImportedSubscriptionKey(row_number=row_number, raw_key=normalized))
                    if len(extracted) > MAX_DELIVERY_ITEMS:
                        raise SubscriptionError("subscription_delivery_item_limit")
    finally:
        workbook.close()
    if not extracted:
        raise SubscriptionError("subscription_delivery_workbook_no_keys")
    return extracted


def validate_imported_subscription_keys(db: Session, admin: User, imported: list[ImportedSubscriptionKey]) -> list[DeliveryValidation]:
    validations: list[DeliveryValidation] = []
    for item in imported:
        key = db.scalar(
            select(PremiumSubscriptionKey).where(
                PremiumSubscriptionKey.tenant_id == admin.tenant_id,
                PremiumSubscriptionKey.key_fingerprint == subscription_key_fingerprint(item.raw_key),
            )
        )
        valid_hash = verify_password(item.raw_key, key.key_hash if key is not None else None)
        if key is None or not valid_hash:
            validations.append(DeliveryValidation(item.row_number, item.raw_key, None, None, None, "invalid"))
        elif not can_manage_institution(db, admin, key.institution_abbr):
            validations.append(DeliveryValidation(item.row_number, item.raw_key, None, None, None, "not_authorized"))
        elif key.status != "issued":
            validations.append(DeliveryValidation(item.row_number, item.raw_key, key.plan_code, key.institution_abbr, key.institution_name_zh, "unavailable"))
        else:
            validations.append(DeliveryValidation(item.row_number, item.raw_key, key.plan_code, key.institution_abbr, key.institution_name_zh, "available"))
    return validations


def _recipient(db: Session, admin: User, user_id) -> User | None:
    return db.scalar(
        select(User)
        .options(selectinload(User.roles), selectinload(User.student_profile), selectinload(User.mentor_profile))
        .where(User.id == user_id, User.tenant_id == admin.tenant_id)
        .with_for_update()
    )


def resolve_delivery_items(
    db: Session,
    admin: User,
    items: list[DeliveryItem],
    *,
    allow_general_delivery: bool = False,
    require_unsubscribed_recipient: bool = False,
) -> list[ResolvedDeliveryItem]:
    if not items or len(items) > MAX_DELIVERY_ITEMS:
        raise SubscriptionError("subscription_delivery_item_limit")
    raw_keys: set[str] = set()
    recipients: set[object] = set()
    resolved: list[ResolvedDeliveryItem] = []
    for item in items:
        raw_key = normalize_subscription_key(item.raw_key)
        if (
            not is_valid_subscription_key_format(raw_key)
            or raw_key in raw_keys
            or (item.recipient_user_id is not None and item.recipient_user_id in recipients)
        ):
            raise SubscriptionError("subscription_delivery_items_invalid")
        raw_keys.add(raw_key)
        if item.recipient_user_id is not None:
            recipients.add(item.recipient_user_id)
        key = db.scalar(
            select(PremiumSubscriptionKey)
            .where(
                PremiumSubscriptionKey.tenant_id == admin.tenant_id,
                PremiumSubscriptionKey.key_fingerprint == subscription_key_fingerprint(raw_key),
            )
            .with_for_update()
        )
        if key is None or not verify_password(raw_key, key.key_hash) or key.status != "issued":
            raise SubscriptionError("subscription_delivery_key_unavailable")
        if not can_manage_institution(db, admin, key.institution_abbr):
            raise SubscriptionError("institution_scope_forbidden")
        if item.recipient_user_id is None:
            if not allow_general_delivery:
                raise SubscriptionError("subscription_delivery_recipient_required")
            resolved.append(ResolvedDeliveryItem(raw_key=raw_key, key=key, recipient=None, general_delivery=True))
            continue

        recipient = _recipient(db, admin, item.recipient_user_id)
        recipient_roles = {role.code for role in recipient.roles} if recipient is not None else set()
        if (
            recipient is None
            or not recipient.is_active
            or not recipient.is_verified
            or not recipient_roles.intersection({"student", "mentor"})
            or user_institution_abbr(recipient) != key.institution_abbr
        ):
            raise SubscriptionError("subscription_delivery_recipient_ineligible")
        if require_unsubscribed_recipient and recipient.id in active_premium_subscription_user_ids(db, [recipient.id]):
            raise SubscriptionError("subscription_delivery_recipient_already_subscribed")
        resolved.append(ResolvedDeliveryItem(raw_key=raw_key, key=key, recipient=recipient))
    return resolved


def activate_subscription_delivery_items(db: Session, admin: User, items: list[DeliveryItem]):
    resolved = resolve_delivery_items(db, admin, items, require_unsubscribed_recipient=True)
    snapshots = []
    now = datetime.now(timezone.utc)
    for item in resolved:
        if item.recipient is None:
            raise SubscriptionError("subscription_delivery_recipient_required")
        snapshots.append(_activate_locked_subscription_key(db, item.recipient, item.key, started_at=now))
    db.commit()
    return snapshots


def build_subscription_delivery_pdf_bundle(db: Session, admin: User, items: list[DeliveryItem]) -> bytes:
    resolved = resolve_delivery_items(db, admin, items, allow_general_delivery=True)
    generated_at = datetime.now(timezone.utc)
    bundle = BytesIO()
    with ZipFile(bundle, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for index, item in enumerate(resolved, start=1):
            document_number = f"AN-SUB-{generated_at.strftime('%Y%m%d')}-{secrets.token_hex(4).upper()}"
            recipient_name = "通用 / GENERAL"
            recipient_username = "不指定 / UNASSIGNED"
            if item.recipient is not None:
                recipient_name = item.recipient.full_name
                recipient_username = item.recipient.username or ""
            pdf = render_subscription_delivery_pdf(
                document_number=document_number,
                recipient_name=recipient_name,
                recipient_username=recipient_username,
                institution_name=item.key.institution_name_zh,
                institution_abbr=item.key.institution_abbr,
                plan_code=item.key.plan_code,
                subscription_key=item.raw_key,
                generated_at=generated_at,
                general_delivery=item.general_delivery,
            )
            filename_seed = "GENERAL" if item.recipient is None else item.recipient.username or str(index)
            safe_username = "".join(character for character in filename_seed if character.isalnum() or character in "_-")[:64]
            archive.writestr(f"AcademicNexus-Subscription-Notice-{index:03d}-{safe_username or index}.pdf", pdf)
    # No database mutation in export mode; locks are released when the request
    # session ends and keys remain redeemable by the intended recipient.
    return bundle.getvalue()
