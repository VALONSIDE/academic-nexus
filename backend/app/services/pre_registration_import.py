"""Validated Excel import and one-time account receipt for partner institutions."""

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import re
from typing import Any
import uuid
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models.pre_registration import PreRegistration, PreRegistrationBatch
from app.models.user import User
from app.services.access_keys import access_key_fingerprint, generate_access_key
from app.services.spreadsheet_safety import safe_spreadsheet_text

TEMPLATE_SHEET = "预注册导入 Pre-registration"
TEMPLATE_HEADERS = (
    "学校英文简称 (School Abbr)",
    "学校中文全称 (Institution Name in Chinese)",
    "学院中文全称 (College Name in Chinese)",
    "用户类型 (Role)",
    "中文真实姓名 (Full Name)",
    "学工号 (Academic ID)",
)
ROLE_ALIASES = {"student": "student", "s": "student", "学生": "student", "mentor": "mentor", "teacher": "mentor", "t": "mentor", "导师": "mentor", "老师": "mentor"}
ABBR_PATTERN = re.compile(r"^[A-Za-z0-9]{2,12}$")
ACADEMIC_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]{2,64}$")
MAX_XLSX_ARCHIVE_ENTRIES = 1000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_IMPORT_ROWS = 1000


class PreRegistrationImportError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))


@dataclass(frozen=True)
class PreparedRow:
    row_number: int
    institution_abbr: str
    institution_name_zh: str
    college_name_zh: str
    role_code: str
    full_name: str
    academic_id: str
    username: str


@dataclass(frozen=True)
class ReceiptRow:
    username: str
    full_name: str
    academic_id: str
    institution_name_zh: str
    college_name_zh: str
    access_key: str


def _cell_text(value: Any) -> str:
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def _username(abbr: str, role: str, academic_id: str) -> str:
    return f"{abbr}_{'S' if role == 'student' else 'T'}{academic_id}".upper()


def _validate_xlsx_archive(content: bytes) -> None:
    """Reject archive bombs before openpyxl reads untrusted workbook XML."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            total_size = sum(entry.file_size for entry in entries)
    except (BadZipFile, OSError) as error:
        raise PreRegistrationImportError(["无法读取 Excel 文件 / Unable to read Excel archive"]) from error
    if not entries or len(entries) > MAX_XLSX_ARCHIVE_ENTRIES or total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise PreRegistrationImportError(["Excel 文件结构不安全或内容过大 / Workbook archive is unsafe or too large"])


def parse_import_workbook(content: bytes) -> list[PreparedRow]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise PreRegistrationImportError([f"无法读取 Excel 文件 / Unable to read Excel: {type(error).__name__}"]) from error
    try:
        return _parse_import_rows(workbook)
    finally:
        workbook.close()


def _parse_import_rows(workbook) -> list[PreparedRow]:
    if TEMPLATE_SHEET not in workbook.sheetnames:
        raise PreRegistrationImportError([f"缺少工作表“{TEMPLATE_SHEET}” / Required worksheet is missing"])
    sheet = workbook[TEMPLATE_SHEET]
    # Ignore attacker-controlled worksheet dimensions and bound both axes.
    sheet.reset_dimensions()
    actual_headers = tuple(_cell_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(TEMPLATE_HEADERS))))
    if actual_headers[:len(TEMPLATE_HEADERS)] != TEMPLATE_HEADERS:
        raise PreRegistrationImportError(["模板列不匹配，请下载并使用最新模板 / Template columns do not match"])
    errors: list[str] = []; prepared: list[PreparedRow] = []; seen: set[str] = set()
    for number, row in enumerate(sheet.iter_rows(min_row=2, max_col=len(TEMPLATE_HEADERS), values_only=True), start=2):
        if number > MAX_IMPORT_ROWS + 1:
            raise PreRegistrationImportError([f"每次最多导入 {MAX_IMPORT_ROWS} 行 / Import row limit exceeded"])
        values = [_cell_text(value) for value in row]
        if not any(values): continue
        abbr, institution, college, raw_role, name, academic_id = values
        abbr, academic_id = abbr.upper(), academic_id.upper()
        role = ROLE_ALIASES.get(raw_role.casefold())
        prefix = f"第 {number} 行 / Row {number}"
        if not ABBR_PATTERN.fullmatch(abbr): errors.append(f"{prefix}: 学校英文简称应为 2-12 位英文字母或数字 / invalid School Abbr")
        elif not 2 <= len(institution) <= 160: errors.append(f"{prefix}: 必须填写学校中文全称 / Institution Name in Chinese is required")
        elif not 2 <= len(college) <= 160: errors.append(f"{prefix}: 必须填写学院中文全称 / College Name in Chinese is required")
        elif role is None: errors.append(f"{prefix}: 用户类型仅支持 student/mentor、S/T、学生/导师 / invalid Role")
        elif not 2 <= len(name) <= 120: errors.append(f"{prefix}: 中文真实姓名长度应为 2-120 / invalid Full Name")
        elif not ACADEMIC_ID_PATTERN.fullmatch(academic_id): errors.append(f"{prefix}: 学工号只能包含字母、数字、下划线或连字符 / invalid Academic ID")
        else:
            username = _username(abbr, role, academic_id)
            if username in seen: errors.append(f"{prefix}: 文件内账户名称重复 / duplicate username: {username}")
            else:
                seen.add(username); prepared.append(PreparedRow(number, abbr, institution, college, role, name, academic_id, username))
    if not prepared: errors.append("模板中没有可导入的数据行 / No importable data rows")
    if errors: raise PreRegistrationImportError(errors)
    return prepared


def _assert_usernames_available(db: Session, tenant_id: uuid.UUID, rows: list[PreparedRow]) -> None:
    usernames = [row.username for row in rows]
    # Login deliberately accepts a username without a tenant selector, so a
    # username must remain globally unique even when more tenants are added.
    pre_registered = set(db.scalars(select(PreRegistration.username).where(PreRegistration.username.in_(usernames))).all())
    activated = set(db.scalars(select(User.username).where(User.username.in_(usernames))).all())
    if unavailable := sorted(pre_registered | activated):
        raise PreRegistrationImportError([f"账户名称已存在且不可再次预注册 / Username already exists: {', '.join(unavailable[:10])}"])


def import_pre_registrations(db: Session, *, tenant_id: uuid.UUID, admin_user_id: uuid.UUID, source_filename: str, content: bytes) -> tuple[PreRegistrationBatch, list[ReceiptRow]]:
    rows = parse_import_workbook(content); _assert_usernames_available(db, tenant_id, rows)
    batch = PreRegistrationBatch(tenant_id=tenant_id, uploaded_by_user_id=admin_user_id, source_filename=source_filename[:255] or "pre-registration.xlsx", row_count=len(rows))
    db.add(batch); db.flush(); receipt_rows: list[ReceiptRow] = []
    for row in rows:
        for _ in range(10):
            key = generate_access_key()
            record = PreRegistration(tenant_id=tenant_id, batch_id=batch.id, username=row.username, role_code=row.role_code, full_name=row.full_name, academic_id=row.academic_id, institution_abbr=row.institution_abbr, institution_name_zh=row.institution_name_zh, college_name_zh=row.college_name_zh, access_key_hash=hash_password(key), access_key_fingerprint=access_key_fingerprint(key), status="issued")
            try:
                with db.begin_nested(): db.add(record); db.flush()
            except IntegrityError: continue
            receipt_rows.append(ReceiptRow(row.username, row.full_name, row.academic_id, row.institution_name_zh, row.college_name_zh, key)); break
        else: raise RuntimeError("Unable to allocate a unique Access Key after 10 attempts")
    batch.receipt_issued_at = datetime.now(timezone.utc); db.commit(); db.refresh(batch)
    return batch, receipt_rows


def _style_header(cells: list[Any]) -> None:
    for cell in cells:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0F172A"); cell.alignment = Alignment(horizontal="center", vertical="center")


def build_import_template() -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = TEMPLATE_SHEET
    sheet.append(list(TEMPLATE_HEADERS)); _style_header(list(sheet[1])); sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEF", (18, 30, 30, 18, 22, 22), strict=True): sheet.column_dimensions[column].width = width
    sheet.column_dimensions["F"].number_format = "@"
    examples = workbook.create_sheet("填写示例 Examples")
    examples.append(list(TEMPLATE_HEADERS) + ["生成账户名称 (Generated Username)"])
    examples.append(["CUC", "中国传媒大学", "新闻学院", "学生", "张同学", "20240001", "CUC_S20240001"])
    examples.append(["CUC", "中国传媒大学", "新闻学院", "导师", "李老师", "T10086", "CUC_TT10086"])
    _style_header(list(examples[1])); examples.freeze_panes = "A2"
    for column, width in zip("ABCDEFG", (18, 30, 30, 18, 22, 22, 30), strict=True): examples.column_dimensions[column].width = width
    examples.column_dimensions["F"].number_format = "@"
    instructions = workbook.create_sheet("填写说明 Instructions")
    instructions.append(["填写说明 / Instructions"])
    for line in ["1. 请勿修改“预注册导入 Pre-registration”工作表名称或表头。", "2. 学校中文全称、学院中文全称、用户类型、姓名和学工号均为必填。", "3. 用户类型可填写 student / mentor、S / T、学生 / 导师。", "4. 学工号必须按文本填写以保留前导零；账户名自动生成。", "5. 系统只在导入回执中提供一次性 Access Key 明文，请妥善发放与保管。"]: instructions.append([line])
    instructions.column_dimensions["A"].width = 100; _style_header([instructions["A1"]])
    stream = BytesIO(); workbook.save(stream); return stream.getvalue()


def build_receipt_workbook(rows: list[ReceiptRow], batch_id: uuid.UUID) -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "账户回执 Account Receipt"
    sheet.append(["账户名称 (Username)", "中文真实姓名 (Full Name)", "学校中文全称", "学院中文全称", "学工号 (Academic ID)", "Access Key"]); _style_header(list(sheet[1]))
    for row in rows:
        sheet.append([
            safe_spreadsheet_text(row.username),
            safe_spreadsheet_text(row.full_name),
            safe_spreadsheet_text(row.institution_name_zh),
            safe_spreadsheet_text(row.college_name_zh),
            safe_spreadsheet_text(row.academic_id),
            safe_spreadsheet_text(row.access_key),
        ])
    for column, width in zip("ABCDEF", (30, 22, 30, 30, 22, 26), strict=True): sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    note = workbook.create_sheet("安全说明 Security Notice"); note.append(["安全说明 / Security Notice"]); _style_header([note["A1"]]); note.append([f"批次 ID / Batch ID: {batch_id}"]); note.append(["请通过受控渠道向本人发放对应行的 Access Key。激活成功后该 Key 自动失效；系统不保存 Key 明文。"]); note.column_dimensions["A"].width = 110
    stream = BytesIO(); workbook.save(stream); return stream.getvalue()
