from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.core.security import verify_password
from app.db.base import Base
from app.models import pre_registration, tenant, user  # noqa: F401
from app.models.pre_registration import PreRegistration, PreRegistrationBatch
from app.models.tenant import Tenant
from app.models.user import User
from app.services.access_keys import access_key_fingerprint, is_valid_access_key_format
from app.services.pre_registration_import import TEMPLATE_HEADERS, TEMPLATE_SHEET, PreRegistrationImportError, build_import_template, build_receipt_workbook, import_pre_registrations, parse_import_workbook
from app.api.v1.endpoints.pre_registrations import _delete_issued_items


def make_workbook(rows: list[tuple[str, str, str, str, str, str]]) -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = TEMPLATE_SHEET; sheet.append(TEMPLATE_HEADERS)
    for row in rows: sheet.append(row)
    stream = BytesIO(); workbook.save(stream); return stream.getvalue()


def test_template_requires_chinese_institution_and_college_names() -> None:
    workbook = load_workbook(BytesIO(build_import_template()), read_only=True, data_only=True)
    assert tuple(cell.value for cell in next(workbook[TEMPLATE_SHEET].iter_rows(min_row=1, max_row=1))) == TEMPLATE_HEADERS
    examples = workbook["填写示例 Examples"]
    assert tuple(cell.value for cell in examples[2][:6]) == ("CUC", "中国传媒大学", "新闻学院", "学生", "张同学", "20240001")
    assert tuple(cell.value for cell in examples[3][:6]) == ("CUC", "中国传媒大学", "新闻学院", "导师", "李老师", "T10086")


def test_partner_import_generates_unique_bound_access_keys_and_identity() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:"); Base.metadata.create_all(engine)
    with Session(engine) as db:
        pilot = Tenant(slug="cuc-pilot", name="CUC Pilot"); db.add(pilot); db.flush()
        admin = User(tenant_id=pilot.id, username="CUC_ADMIN", full_name="Admin", password_hash="hash"); db.add(admin); db.flush()
        batch, receipt = import_pre_registrations(db, tenant_id=pilot.id, admin_user_id=admin.id, source_filename="cuc.xlsx", content=make_workbook([("CUC", "中国传媒大学", "新闻学院", "学生", "张同学", "20240001"), ("CUC", "中国传媒大学", "新闻学院", "mentor", "李老师", "T10086")]))
        assert batch.row_count == 2 and [row.username for row in receipt] == ["CUC_S20240001", "CUC_TT10086"]
        assert all(is_valid_access_key_format(row.access_key) for row in receipt)
        records = db.scalars(select(PreRegistration).order_by(PreRegistration.username)).all(); assert records[0].institution_name_zh == "中国传媒大学" and records[0].college_name_zh == "新闻学院"
        assert verify_password(receipt[0].access_key, {item.username: item for item in records}[receipt[0].username].access_key_hash)
        assert {item.access_key_fingerprint for item in records} == {access_key_fingerprint(row.access_key) for row in receipt}
        sheet = load_workbook(BytesIO(build_receipt_workbook(receipt, batch.id)), data_only=True)["账户回执 Account Receipt"]
        assert sheet.max_row == 3 and sheet["F2"].value == receipt[0].access_key


def test_import_rejects_missing_identity_or_duplicate_username() -> None:
    missing_college = make_workbook([("CUC", "中国传媒大学", "", "student", "张同学", "20240001")])
    with pytest.raises(PreRegistrationImportError, match="College Name"):
        parse_import_workbook(missing_college)
    duplicate = make_workbook([("CUC", "中国传媒大学", "新闻学院", "student", "张同学", "20240001"), ("CUC", "中国传媒大学", "新闻学院", "S", "王同学", "20240001")])
    with pytest.raises(PreRegistrationImportError, match="duplicate username"):
        parse_import_workbook(duplicate)


def test_admin_can_delete_only_unactivated_pre_registration_accounts() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:"); Base.metadata.create_all(engine)
    with Session(engine) as db:
        pilot = Tenant(slug="delete-test", name="Delete Test"); db.add(pilot); db.flush()
        admin = User(tenant_id=pilot.id, username="TEST_ADMIN", full_name="Admin", password_hash="hash"); db.add(admin); db.flush()
        batch = PreRegistrationBatch(tenant_id=pilot.id, uploaded_by_user_id=admin.id, source_filename="test.xlsx", row_count=2); db.add(batch); db.flush()
        issued = PreRegistration(tenant_id=pilot.id, batch_id=batch.id, username="TEST_S1", role_code="student", full_name="Student", academic_id="S1", institution_abbr="TEST", institution_name_zh="Test University", college_name_zh="Test College", access_key_hash="hash", access_key_fingerprint="a" * 64, status="issued")
        activated = PreRegistration(tenant_id=pilot.id, batch_id=batch.id, username="TEST_T1", role_code="mentor", full_name="Mentor", academic_id="T1", institution_abbr="TEST", institution_name_zh="Test University", college_name_zh="Test College", access_key_hash="hash", access_key_fingerprint="b" * 64, status="activated")
        db.add_all([issued, activated]); db.flush()

        assert _delete_issued_items(db, admin, [issued.id]) == 1
        assert db.get(PreRegistration, issued.id) is None
        with pytest.raises(HTTPException) as error:
            _delete_issued_items(db, admin, [activated.id])
        assert error.value.status_code == 409
        assert error.value.detail["code"] == "activated_pre_registration_cannot_be_deleted"
