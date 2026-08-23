from datetime import datetime, timedelta, timezone
from io import BytesIO
import re
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.models import ai, resource, selection, subscription, tenant, user  # noqa: F401
from app.models.ai import AiUserSubscription
from app.models.subscription import PremiumSubscriptionKey
from app.models.tenant import Tenant
from app.models.user import Role, StudentProfile, User
from app.services.ai.assistant import create_conversation, delete_conversation, list_conversations
from app.api.v1.endpoints.users import _load_managed_user
from app.api.v1.endpoints.dashboard import read_dashboard
from app.services.subscriptions import (
    SubscriptionError,
    activate_subscription_key,
    assign_institution_admin,
    build_subscription_key_receipt,
    issue_subscription_key,
    issue_subscription_keys,
    refresh_subscription,
    revoke_subscription_key,
    set_institution_allocation,
)
from app.services.subscription_delivery import (
    DeliveryItem,
    activate_subscription_delivery_items,
    build_subscription_delivery_pdf_bundle,
    extract_subscription_keys_from_workbook,
    validate_imported_subscription_keys,
)


def _role(db: Session, organization: Tenant, code: str) -> Role:
    item = Role(tenant_id=organization.id, code=code, name_zh=code, name_en=code)
    db.add(item)
    db.flush()
    return item


def _user(db: Session, organization: Tenant, username: str, roles: list[Role], *, created_at: datetime) -> User:
    item = User(
        tenant_id=organization.id,
        username=username,
        full_name=username,
        password_hash="hash",
        created_at=created_at,
        is_active=True,
        is_verified=True,
    )
    item.roles.extend(roles)
    db.add(item)
    db.flush()
    return item


def test_conversations_are_owned_deletable_and_capped_at_one_hundred() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="history", name="History")
        db.add(organization)
        db.flush()
        student_role = _role(db, organization, "student")
        account = _user(db, organization, "HIST_S001", [student_role], created_at=datetime(2026, 1, 1, tzinfo=timezone.utc))
        other = _user(db, organization, "HIST_S002", [student_role], created_at=datetime(2026, 1, 1, tzinfo=timezone.utc))
        db.commit()

        first = create_conversation(db, account, topic="academic_planning", title="first")
        for number in range(100):
            create_conversation(db, account, topic="academic_planning", title=f"conversation {number}")
        kept = list_conversations(db, account)
        assert len(kept) == 100
        assert all(item.id != first.id for item in kept)

        other_conversation = create_conversation(db, other, topic="academic_planning", title="other")
        delete_conversation(db, other_conversation)
        assert len(list_conversations(db, other)) == 0
        assert len(list_conversations(db, account)) == 100


def test_premium_key_is_institution_bound_reclaimable_before_use_and_resets_to_basic() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="subscription", name="Subscription")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        institution_role = _role(db, organization, "institution_admin")
        student_role = _role(db, organization, "student")
        started_at = datetime(2026, 1, 15, 9, tzinfo=timezone.utc)
        super_admin = _user(db, organization, "SUB_A001", [super_role], created_at=started_at)
        manager = _user(db, organization, "SUB_A002", [institution_role], created_at=started_at)
        student = _user(db, organization, "SUB_S003", [student_role], created_at=started_at)
        other_student = _user(db, organization, "SUB_S004", [student_role], created_at=started_at)
        student.student_profile = StudentProfile(user_id=student.id, institution_abbr="CUC", university="中国传媒大学")
        other_student.student_profile = StudentProfile(user_id=other_student.id, institution_abbr="THU", university="清华大学")
        manager.student_profile = StudentProfile(user_id=manager.id, institution_abbr="CUC", university="中国传媒大学")
        db.commit()

        set_institution_allocation(db, super_admin, institution_abbr="CUC", pro_credits=2, ultra_credits=0, max_credits=0)
        assign_institution_admin(db, super_admin, user_id=manager.id, institution_abbr="CUC")
        issued, plaintext = issue_subscription_key(db, manager, institution_abbr="CUC", plan_code="pro")
        assert re.fullmatch(r"[A-Z0-9]{2}-P-[A-Z0-9]{4}-[A-Z0-9]{3}", plaintext)
        reclaimed = revoke_subscription_key(db, manager, issued.id)
        assert reclaimed.status == "revoked"

        active_key, active_plaintext = issue_subscription_key(db, manager, institution_abbr="CUC", plan_code="pro")
        with pytest.raises(SubscriptionError, match="subscription_key_institution_mismatch"):
            activate_subscription_key(db, other_student, raw_key=active_plaintext)
        assert db.get(PremiumSubscriptionKey, active_key.id).status == "issued"
        activated = activate_subscription_key(db, student, raw_key=active_plaintext)
        assert activated.plan_code == "pro"
        assert activated.credit_limit == 50
        assert activated.credit_balance == 50
        assert db.get(PremiumSubscriptionKey, active_key.id).status == "activated"
        with pytest.raises(SubscriptionError, match="subscription_key_not_reclaimable"):
            revoke_subscription_key(db, manager, active_key.id)

        current = db.scalar(select(AiUserSubscription).where(AiUserSubscription.user_id == student.id))
        assert current is not None
        refreshed = refresh_subscription(db, student, now=current.cycle_ends_at + timedelta(seconds=1))
        assert refreshed.plan_code == "basic"
        assert refreshed.credit_limit == 10


def test_institution_administrator_cannot_manage_users_outside_assigned_institution() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="scope", name="Scope")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        institution_role = _role(db, organization, "institution_admin")
        student_role = _role(db, organization, "student")
        created_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
        super_admin = _user(db, organization, "SCOPE_A001", [super_role], created_at=created_at)
        manager = _user(db, organization, "SCOPE_A002", [institution_role], created_at=created_at)
        cuc_student = _user(db, organization, "SCOPE_S003", [student_role], created_at=created_at)
        other_student = _user(db, organization, "SCOPE_S004", [student_role], created_at=created_at)
        cuc_student.student_profile = StudentProfile(user_id=cuc_student.id, institution_abbr="CUC", university="中国传媒大学")
        other_student.student_profile = StudentProfile(user_id=other_student.id, institution_abbr="THU", university="清华大学")
        manager.student_profile = StudentProfile(user_id=manager.id, institution_abbr="CUC", university="中国传媒大学")
        db.commit()

        assign_institution_admin(db, super_admin, user_id=manager.id, institution_abbr="CUC")
        assert _load_managed_user(db, manager, cuc_student.id).id == cuc_student.id
        with pytest.raises(HTTPException) as forbidden:
            _load_managed_user(db, manager, other_student.id)
        assert forbidden.value.status_code == 404
        assert _load_managed_user(db, super_admin, other_student.id).id == other_student.id


def test_batch_subscription_keys_are_atomic_inventory_checked_and_exported_once() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="batch", name="Batch")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        student_role = _role(db, organization, "student")
        started_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
        super_admin = _user(db, organization, "BATCH_A001", [super_role], created_at=started_at)
        institution_student = _user(db, organization, "BATCH_S002", [student_role], created_at=started_at)
        institution_student.student_profile = StudentProfile(user_id=institution_student.id, institution_abbr="CUC", university="中国传媒大学")
        db.commit()
        allocation = set_institution_allocation(
            db,
            super_admin,
            institution_abbr="CUC",
            pro_credits=2,
            ultra_credits=0,
            max_credits=0,
        )

        issued = issue_subscription_keys(db, super_admin, institution_abbr="CUC", plan_code="pro", quantity=2)
        assert len(issued) == 2
        assert all(re.fullmatch(r"[A-Z0-9]{2}-P-[A-Z0-9]{4}-[A-Z0-9]{3}", item.plaintext) for item in issued)
        assert all(item.plaintext not in item.key.key_hash for item in issued)
        db.refresh(allocation)
        assert allocation.pro_credits == 0
        with pytest.raises(SubscriptionError, match="subscription_inventory_exhausted"):
            issue_subscription_keys(db, super_admin, institution_abbr="CUC", plan_code="pro", quantity=1)

        workbook = load_workbook(BytesIO(build_subscription_key_receipt(issued)), data_only=True)
        sheet = workbook["订阅 Key 回执 Subscription Keys"]
        assert sheet.max_row == 3
        assert sheet.cell(2, 1).value == issued[0].plaintext

        issued[0].key.institution_name_zh = "=FORMULA()"
        protected = load_workbook(BytesIO(build_subscription_key_receipt([issued[0]])), data_only=False)
        assert protected["订阅 Key 回执 Subscription Keys"]["C2"].value == "'=FORMULA()"


def test_institution_administrator_dashboard_is_scoped_to_their_institution() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="dashboard-scope", name="Dashboard scope")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        institution_role = _role(db, organization, "institution_admin")
        student_role = _role(db, organization, "student")
        created_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
        super_admin = _user(db, organization, "DASH_A001", [super_role], created_at=created_at)
        manager = _user(db, organization, "DASH_A002", [institution_role], created_at=created_at)
        cuc_student = _user(db, organization, "DASH_S003", [student_role], created_at=created_at)
        other_student = _user(db, organization, "DASH_S004", [student_role], created_at=created_at)
        manager.student_profile = StudentProfile(user_id=manager.id, institution_abbr="CUC", university="中国传媒大学")
        cuc_student.student_profile = StudentProfile(user_id=cuc_student.id, institution_abbr="CUC", university="中国传媒大学")
        other_student.student_profile = StudentProfile(user_id=other_student.id, institution_abbr="THU", university="清华大学")
        db.commit()

        assign_institution_admin(db, super_admin, user_id=manager.id, institution_abbr="CUC")
        dashboard = read_dashboard(manager, db)
        metrics = {item.code: item.value for item in dashboard.metrics}
        assert metrics["students_total"] == 1
        assert metrics["mentors_total"] == 0


def test_offline_key_receipt_is_validated_in_memory_then_can_directly_activate() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="offline-delivery", name="Offline delivery")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        student_role = _role(db, organization, "student")
        started_at = datetime(2026, 8, 23, tzinfo=timezone.utc)
        admin = _user(db, organization, "DELIVERY_A001", [super_role], created_at=started_at)
        recipient = _user(db, organization, "DELIVERY_S002", [student_role], created_at=started_at)
        recipient.student_profile = StudentProfile(user_id=recipient.id, institution_abbr="CUC", university="China Communication University")
        db.commit()
        set_institution_allocation(db, admin, institution_abbr="CUC", pro_credits=1, ultra_credits=0, max_credits=0)
        issued, plaintext = issue_subscription_key(db, admin, institution_abbr="CUC", plan_code="pro")

        workbook = Workbook()
        workbook.active.append(["Subscription Key"])
        workbook.active.append([plaintext])
        source = BytesIO()
        workbook.save(source)
        imported = extract_subscription_keys_from_workbook(source.getvalue())
        validation = validate_imported_subscription_keys(db, admin, imported)
        assert len(validation) == 1
        assert validation[0].status == "available"
        assert validation[0].raw_key == plaintext

        bundle = build_subscription_delivery_pdf_bundle(db, admin, [DeliveryItem(raw_key=plaintext, recipient_user_id=recipient.id)])
        with ZipFile(BytesIO(bundle)) as archive:
            names = archive.namelist()
            assert len(names) == 1 and names[0].endswith(".pdf")
            assert archive.read(names[0]).startswith(b"%PDF")
        assert db.get(PremiumSubscriptionKey, issued.id).status == "issued"

        snapshots = activate_subscription_delivery_items(db, admin, [DeliveryItem(raw_key=plaintext, recipient_user_id=recipient.id)])
        assert snapshots[0].plan_code == "pro"
        assert db.get(PremiumSubscriptionKey, issued.id).status == "activated"
        assert plaintext not in db.get(PremiumSubscriptionKey, issued.id).key_hash


def test_offline_direct_activation_rejects_active_premium_and_allows_general_pdf() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        organization = Tenant(slug="offline-general", name="Offline general")
        db.add(organization)
        db.flush()
        super_role = _role(db, organization, "super_admin")
        student_role = _role(db, organization, "student")
        started_at = datetime(2026, 8, 23, tzinfo=timezone.utc)
        admin = _user(db, organization, "GENERAL_A001", [super_role], created_at=started_at)
        recipient = _user(db, organization, "GENERAL_S002", [student_role], created_at=started_at)
        recipient.student_profile = StudentProfile(user_id=recipient.id, institution_abbr="CUC", university="China Communication University")
        db.commit()

        set_institution_allocation(db, admin, institution_abbr="CUC", pro_credits=2, ultra_credits=0, max_credits=0)
        _, first_plaintext = issue_subscription_key(db, admin, institution_abbr="CUC", plan_code="pro")
        activate_subscription_key(db, recipient, raw_key=first_plaintext)
        next_key, next_plaintext = issue_subscription_key(db, admin, institution_abbr="CUC", plan_code="pro")

        with pytest.raises(SubscriptionError, match="subscription_delivery_recipient_already_subscribed"):
            activate_subscription_delivery_items(db, admin, [DeliveryItem(raw_key=next_plaintext, recipient_user_id=recipient.id)])
        db.rollback()
        assert db.get(PremiumSubscriptionKey, next_key.id).status == "issued"

        bundle = build_subscription_delivery_pdf_bundle(db, admin, [DeliveryItem(raw_key=next_plaintext, recipient_user_id=None)])
        with ZipFile(BytesIO(bundle)) as archive:
            document = archive.read(archive.namelist()[0])
            assert document.startswith(b"%PDF")
            assert len(document) > 4_000
        assert db.get(PremiumSubscriptionKey, next_key.id).status == "issued"
